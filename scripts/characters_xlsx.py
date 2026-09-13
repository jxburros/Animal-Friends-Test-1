#!/usr/bin/env python3
"""Bind the two character CSVs into one .xlsx workbook.

    node scripts/characters.mjs && python3 scripts/characters_xlsx.py

Writes docs/character_versions.xlsx with two sheets — "Characters" (one row per named Character)
and "Versions" (one row per printed version) — each with a frozen, filterable header row and
columns sized to fit. The CSVs written by scripts/characters.mjs are the source of truth; this
script only formats them, so it needs openpyxl (pip install openpyxl) and nothing else.
"""
import csv
import pathlib
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - the CSVs are still the deliverable without it
    sys.exit("openpyxl is not installed (pip install openpyxl). The CSVs in docs/ are already written.")

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
OUT = DOCS / "character_versions.xlsx"
SHEETS = [("Characters", DOCS / "characters.csv"), ("Versions", DOCS / "character_versions.csv")]

HEADER_FILL = PatternFill("solid", fgColor="F3E3C2")
MAX_WIDTH = 60


def add_sheet(book, title, path, first):
    if not path.exists():
        sys.exit(f"{path.relative_to(ROOT)} is missing — run: node scripts/characters.mjs")
    sheet = book.active if first else book.create_sheet()
    sheet.title = title
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, _ in enumerate(rows[0], start=1):
        width = max(len(str(row[i - 1])) for row in rows if len(row) >= i) + 2
        sheet.column_dimensions[get_column_letter(i)].width = min(width, MAX_WIDTH)


def main():
    book = Workbook()
    for i, (title, path) in enumerate(SHEETS):
        add_sheet(book, title, path, i == 0)
    book.save(OUT)
    print(f"{OUT.relative_to(ROOT)} — {len(SHEETS)} sheets")


if __name__ == "__main__":
    main()
